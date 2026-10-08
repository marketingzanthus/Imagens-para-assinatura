"""Real PostgreSQL integration checks; each run uses a disposable schema."""
import io
import os
from pathlib import Path
import secrets
import sqlite3
import sys
import tempfile
import unittest

import psycopg
from psycopg import sql
from psycopg.conninfo import make_conninfo
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
BASE_URL = os.environ['TEST_DATABASE_URL']
SCHEMA = 'test_' + secrets.token_hex(8)
with psycopg.connect(BASE_URL, autocommit=True) as admin_connection:
    admin_connection.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(SCHEMA)))
os.environ['DATABASE_URL'] = make_conninfo(BASE_URL, options=f'-c search_path={SCHEMA}')
os.environ['ADMIN_PASSWORD'] = secrets.token_urlsafe(32)
os.environ['SECRET_KEY'] = secrets.token_hex(32)
os.environ['RENDER'] = 'true'

class Integration(unittest.TestCase):
    def test_import_and_flows(self):
        from migrate_sqlite import migrate
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'source.db'
            from database import TABLES
            connection = sqlite3.connect(source)
            for table, columns in TABLES.items():
                connection.execute(f'CREATE TABLE {table}(id INTEGER PRIMARY KEY,{columns})')
            # Sparse IDs and exact text/hash preservation across the migration.
            from werkzeug.security import generate_password_hash
            password_hash = generate_password_hash(os.environ['ADMIN_PASSWORD'])
            connection.execute("INSERT INTO users(id,username,password,name,role) VALUES(9,'admin',?,'Administrador','admin')", (password_hash,))
            connection.execute("INSERT INTO campaigns(id,name,title,prize,terms,active,created_at) VALUES(12,'Migração','Teste','Kit','Termos',1,'2026-10-08 20:40:01')")
            connection.execute("INSERT INTO participants(id,campaign_id,name,email,company) VALUES(40,12,'Preservado','old@example.invalid','Empresa')")
            connection.commit(); connection.close()
            counts = migrate(source, Path(directory)/'backup.db')
            self.assertEqual(counts['campaigns'],1)
            self.assertEqual(counts['participants'],1)
            with self.assertRaises(ValueError):
                migrate(source, Path(directory)/'second-backup.db')
        import app as module
        module.app.config['TESTING']=True
        client=module.app.test_client()
        self.assertEqual(client.get('/health').json['database'],'postgresql')
        self.assertEqual(client.post('/admin/login',data={'username':'admin','password':os.environ['ADMIN_PASSWORD']}).status_code,302)
        self.assertEqual(client.get('/admin').status_code,200)
        def post(path, **values):
            with client.session_transaction() as session:
                session['csrf']='test'
            return client.post(path,data={'csrf':'test',**values})
        response=post('/admin/campaign',name='Nova',title='Nova',prize='Kit',terms='Termos',active='on')
        cid=int(response.location.split('=')[1]); self.assertEqual(cid,13)
        values=dict(name='Teste',whatsapp='11999999999',email='new@example.invalid',company='Empresa',consent='on')
        self.assertEqual(post(f'/cadastro/{cid}',**values).location,f'/sucesso/{cid}')
        post(f'/cadastro/{cid}',**values)
        connection=module.db()
        self.assertEqual(connection.execute('SELECT count(*) FROM participants WHERE campaign_id=?',(cid,)).fetchone()[0],1)
        participant=connection.execute('SELECT * FROM participants WHERE campaign_id=?',(cid,)).fetchone()
        self.assertEqual(participant['id'],41)
        self.assertEqual(connection.execute('SELECT password FROM users WHERE id=9').fetchone()[0],password_hash)
        connection.close()
        self.assertEqual(post('/admin/draw',cid=str(cid),count='1',seconds='3').status_code,302)
        self.assertEqual(client.get('/admin/draw-screen?ids=41').status_code,200)
        wb=load_workbook(io.BytesIO(client.get(f'/admin/export?cid={cid}').data))
        self.assertEqual(wb.active.cell(2,5).value,'Sim')
        post('/admin/p/41/block')
        connection=module.db()
        self.assertEqual(connection.execute('SELECT blocked FROM participants WHERE id=41').fetchone()[0],1)
        connection.close()
        self.assertEqual(client.get(f'/qr/{cid}.png').content_type,'image/png')
        self.assertEqual(client.post('/admin/campaign',data={'image':(io.BytesIO(b'img'),'test.png')}).status_code,302)
        post('/admin/user',name='Operador',username='operator',password='test-only-password')
        post('/admin/user',name='Operador',username='operator',password='test-only-password')
        # Restart/bootstrap must neither duplicate records nor reset credentials.
        module.init()
        connection=module.db()
        self.assertEqual(connection.execute('SELECT count(*) FROM campaigns').fetchone()[0],2)
        self.assertEqual(connection.execute('SELECT count(*) FROM users').fetchone()[0],2)
        connection.close()

try:
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(Integration))
finally:
    with psycopg.connect(BASE_URL,autocommit=True) as admin_connection:
        admin_connection.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(SCHEMA)))
raise SystemExit(0 if result.wasSuccessful() else 1)
